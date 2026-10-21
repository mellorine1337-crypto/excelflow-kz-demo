"""Data ingestion, validation and reporting for ExcelFlow KZ."""
from __future__ import annotations
import io, re, unicodedata
from pathlib import Path
from typing import BinaryIO, Mapping
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_FILE_SIZE, MAX_ROWS, MAX_COLUMNS = 10*1024*1024, 200_000, 100
REQUIRED_FIELDS = ("date", "product", "quantity", "price", "cost")
OPTIONAL_FIELDS = ("category",)
FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS
FIELD_LABELS = {"date":"Дата продажи","product":"Товар","quantity":"Количество","price":"Цена за единицу","cost":"Себестоимость за единицу","category":"Категория (необязательно)"}
ALIASES = {
 "date":{"date","дата","дата продажи","дата заказа","день","sales date","order date","дата документа","дата чека","дата и время"},
 "product":{"product","товар","продукт","наименование","номенклатура","название товара","item","sku name","наименование позиции","блюдо","позиция"},
 "quantity":{"quantity","количество","кол во","колво","шт","штук","qty","units","продано","количество товара","количество позиций"},
 "price":{"price","цена","цена продажи","стоимость","цена за единицу","unit price","sale price","цена реализации","сумма за единицу"},
 "cost":{"cost","себестоимость","закупочная цена","закупка","себес","unit cost","cogs","себестоимость единицы"},
 "category":{"category","категория","группа","группа товара","тип товара","product category"},
}

def _clean_name(value):
    text=unicodedata.normalize("NFKC",str(value)).strip().lower().replace("ё","е")
    return re.sub(r"[^\w]+"," ",text,flags=re.UNICODE).strip()

def _file_bytes(uploaded: BinaryIO):
    name=Path(getattr(uploaded,"name","")).name
    raw=uploaded.getvalue() if hasattr(uploaded,"getvalue") else uploaded.read()
    if not raw: raise ValueError("Файл пустой.")
    if len(raw)>MAX_FILE_SIZE: raise ValueError("Размер файла превышает 10 МБ.")
    suffix=Path(name).suffix.lower()
    if suffix not in {".xlsx",".csv"}: raise ValueError("Поддерживаются только файлы XLSX и CSV.")
    return raw,suffix

def excel_sheets(uploaded):
    raw,suffix=_file_bytes(uploaded)
    if suffix!=".xlsx": return []
    try: return pd.ExcelFile(io.BytesIO(raw),engine="openpyxl").sheet_names
    except Exception as exc: raise ValueError("Не удалось открыть Excel-файл. Проверьте, что он не повреждён и не защищён паролем.") from exc

def read_table(uploaded, sheet_name=0):
    raw,suffix=_file_bytes(uploaded)
    try:
        if suffix==".csv":
            try: df=pd.read_csv(io.BytesIO(raw),sep=None,engine="python",encoding="utf-8-sig")
            except UnicodeDecodeError: df=pd.read_csv(io.BytesIO(raw),sep=None,engine="python",encoding="cp1251")
        else: df=pd.read_excel(io.BytesIO(raw),sheet_name=sheet_name,engine="openpyxl")
    except ValueError: raise
    except Exception as exc: raise ValueError("Не удалось прочитать файл. Проверьте формат, разделитель и заголовки таблицы.") from exc
    if df.empty: raise ValueError("В выбранном листе нет строк с данными.")
    if len(df)>MAX_ROWS: raise ValueError(f"Слишком много строк: {len(df):,}. Максимум — {MAX_ROWS:,}.")
    if len(df.columns)>MAX_COLUMNS: raise ValueError(f"Слишком много колонок: {len(df.columns)}. Максимум — {MAX_COLUMNS}.")
    names=[str(c).strip() for c in df.columns]
    if any(not n or n.lower().startswith("unnamed:") for n in names): raise ValueError("У таблицы есть пустые заголовки колонок. Заполните первую строку заголовков.")
    if len({_clean_name(n) for n in names})!=len(names): raise ValueError("В таблице есть повторяющиеся названия колонок.")
    df.columns=names
    return df

def suggest_mapping(columns):
    result={}; cleaned={str(c):_clean_name(c) for c in columns}
    for field in FIELDS:
        aliases={_clean_name(a) for a in ALIASES[field]}
        exact=[c for c,n in cleaned.items() if n in aliases]
        partial=[c for c,n in cleaned.items() if any(a in n or n in a for a in aliases)]
        result[field]=exact[0] if exact else (partial[0] if field!="category" and len(partial)==1 else None)
    return result

def _number_series(series):
    if pd.api.types.is_numeric_dtype(series): return pd.to_numeric(series,errors="coerce")
    normalized=(series.astype("string").str.replace("\u00a0","",regex=False).str.replace(" ","",regex=False).str.replace("₸","",regex=False).str.replace("тг","",case=False,regex=False).str.replace(",",".",regex=False))
    return pd.to_numeric(normalized,errors="coerce")

def prepare_sales(raw_df: pd.DataFrame, mapping: Mapping):
    missing=[FIELD_LABELS[f] for f in REQUIRED_FIELDS if not mapping.get(f)]
    if missing: raise ValueError("Сопоставьте обязательные поля: "+", ".join(missing)+".")
    selected=[str(mapping[f]) for f in REQUIRED_FIELDS]
    if mapping.get("category"): selected.append(str(mapping["category"]))
    if len(set(selected))!=len(selected): raise ValueError("Одну колонку нельзя использовать сразу для нескольких полей.")
    unknown=[c for c in selected if c not in raw_df.columns]
    if unknown: raise ValueError("Не найдены выбранные колонки: "+", ".join(unknown)+".")
    df=raw_df[selected].copy(); df.columns=list(REQUIRED_FIELDS)+(["category"] if mapping.get("category") else [])
    if "category" not in df: df["category"]="Без категории"
    df["date"]=pd.to_datetime(df["date"],errors="coerce",dayfirst=True)
    df["product"]=df["product"].astype("string").str.strip()
    df["category"]=df["category"].astype("string").str.strip().fillna("Без категории").replace("","Без категории")
    for col in ("quantity","price","cost"): df[col]=_number_series(df[col])
    masks={"дата":df.date.isna(),"товар":df["product"].isna()|df["product"].eq(""),"числа":df[["quantity","price","cost"]].isna().any(axis=1),"отрицательные значения":(df[["quantity","price","cost"]]<0).any(axis=1)}
    invalid=pd.concat(masks.values(),axis=1).any(axis=1)
    if invalid.any():
        details="; ".join(f"{label}: {int(mask.sum())}" for label,mask in masks.items() if mask.any())
        rows=", ".join(str(i+2) for i in df.index[invalid][:8])
        raise ValueError(f"Некорректных строк: {int(invalid.sum())} ({details}). Проверьте строки Excel: {rows}.")
    if (df.quantity==0).all(): raise ValueError("Во всех строках количество равно нулю.")
    df["revenue"]=df.quantity*df.price; df["gross_profit"]=df.quantity*(df.price-df.cost)
    df["margin_pct"]=(df.gross_profit/df.revenue.where(df.revenue!=0)*100).fillna(0)
    return df.sort_values("date").reset_index(drop=True)

def load_sales(uploaded,sheet_name=0,mapping=None):
    raw=read_table(uploaded,sheet_name); return prepare_sales(raw,mapping or suggest_mapping(raw.columns))

def summarize(df):
    revenue=float(df.revenue.sum()); profit=float(df.gross_profit.sum())
    return {"revenue":revenue,"gross_profit":profit,"margin_pct":profit/revenue*100 if revenue else 0.,"units":float(df.quantity.sum()),"orders":float(len(df)),"avg_check":revenue/len(df) if len(df) else 0.}

def period_comparison(all_df,start,end):
    start,end=pd.Timestamp(start).normalize(),pd.Timestamp(end).normalize()
    if end<start: raise ValueError("Дата окончания периода раньше даты начала.")
    days=int((end-start).days)+1; prev_end=start-pd.Timedelta(1,unit="D"); prev_start=prev_end-pd.Timedelta(days-1,unit="D")
    current=all_df[all_df.date.dt.normalize().between(start,end)].copy(); previous=all_df[all_df.date.dt.normalize().between(prev_start,prev_end)].copy()
    cs,ps=summarize(current),summarize(previous); change={}
    for key in ("revenue","gross_profit","units","orders","avg_check"):
        change[key]=(cs[key]-ps[key])/abs(ps[key])*100 if ps[key] else None
    change["margin_pct"]=cs["margin_pct"]-ps["margin_pct"] if len(previous) else None
    return current,previous,change

def generate_insights(current,previous=None):
    if current.empty:return []
    s=summarize(current); bp=current.groupby("product",observed=True)[["revenue","gross_profit"]].sum(); top=bp.revenue.idxmax(); share=bp.loc[top,"revenue"]/s["revenue"]*100 if s["revenue"] else 0
    out=[f"Лидер по выручке — {top}: {share:.1f}% выручки выбранного периода."]
    loss=bp[bp.gross_profit<0].sort_values("gross_profit")
    if not loss.empty: out.append(f"{len(loss)} товар(а) продаются с отрицательной валовой прибылью; первым проверьте «{loss.index[0]}».")
    elif s["margin_pct"]<20: out.append(f"Валовая маржа {s['margin_pct']:.1f}% — стоит проверить цены и закупочную стоимость.")
    else: out.append(f"Валовая маржа периода — {s['margin_pct']:.1f}%; убыточных товаров нет.")
    if previous is not None and not previous.empty:
        old=summarize(previous)["revenue"]; delta=(s["revenue"]-old)/abs(old)*100 if old else None
        if delta is not None: out.append(f"Выручка {'выросла' if delta>=0 else 'снизилась'} на {abs(delta):.1f}% относительно предыдущего периода такой же длины.")
    daily=current.groupby(current.date.dt.date).revenue.sum()
    if len(daily)>1: out.append(f"Лучший день — {daily.idxmax():%d.%m.%Y}: {daily.max():,.0f} ₸ выручки.")
    return out

def product_abc(df):
    """Classify products by cumulative revenue share: A <=80%, B <=95%, C remainder."""
    grouped=df.groupby("product",as_index=False,observed=True)[["revenue","gross_profit","quantity"]].sum().sort_values("revenue",ascending=False)
    total=grouped.revenue.sum()
    grouped["revenue_share_pct"]=(grouped.revenue/total*100 if total else 0)
    grouped["cumulative_share_pct"]=grouped.revenue_share_pct.cumsum()
    grouped["abc_class"]=grouped.cumulative_share_pct.map(lambda x:"A" if x<=80 else ("B" if x<=95 else "C"))
    if len(grouped): grouped.loc[grouped.index[0],"abc_class"]="A"
    return grouped

def export_report(df,previous=None):
    if df.empty: raise ValueError("Нечего экспортировать: выбранный период пуст.")
    output=io.BytesIO(); bp=df.groupby(["category","product"],as_index=False,observed=True)[["revenue","gross_profit","quantity"]].sum().sort_values("revenue",ascending=False); bp["margin_pct"]=(bp.gross_profit/bp.revenue.where(bp.revenue!=0)*100).fillna(0)
    bd=df.assign(date=df.date.dt.normalize()).groupby("date",as_index=False)[["revenue","gross_profit","quantity"]].sum()
    totals=pd.DataFrame([summarize(df)],index=["Выбранный период"])
    if previous is not None and not previous.empty: totals=pd.concat([totals,pd.DataFrame([summarize(previous)],index=["Предыдущий период"])])
    sales_export=df.rename(columns={"date":"Дата","category":"Категория","product":"Товар","quantity":"Количество","price":"Цена за единицу","cost":"Себестоимость за единицу","revenue":"Выручка","gross_profit":"Валовая прибыль","margin_pct":"Маржа, %"})
    products_export=bp.rename(columns={"category":"Категория","product":"Товар","revenue":"Выручка","gross_profit":"Валовая прибыль","quantity":"Количество","margin_pct":"Маржа, %"})
    daily_export=bd.rename(columns={"date":"Дата","revenue":"Выручка","gross_profit":"Валовая прибыль","quantity":"Количество"})
    totals_export=totals.rename(columns={"revenue":"Выручка","gross_profit":"Валовая прибыль","margin_pct":"Маржа, %","units":"Количество","orders":"Строк продаж","avg_check":"Средняя сумма строки"})
    with pd.ExcelWriter(output,engine="openpyxl") as writer:
        sales_export.to_excel(writer,sheet_name="Продажи",index=False); products_export.to_excel(writer,sheet_name="По товарам",index=False); daily_export.to_excel(writer,sheet_name="По дням",index=False); totals_export.rename_axis("Период").reset_index().to_excel(writer,sheet_name="Итоги",index=False); pd.DataFrame({"Аналитический вывод":generate_insights(df,previous)}).to_excel(writer,sheet_name="Выводы",index=False)
        fill=PatternFill("solid",fgColor="17324D")
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.sheet_view.showGridLines=False
            for cell in ws[1]: cell.fill=fill; cell.font=Font(color="FFFFFF",bold=True); cell.alignment=Alignment(vertical="center")
            ws.row_dimensions[1].height=24
            for idx,cells in enumerate(ws.columns,1):
                vals=[str(c.value or "") for c in list(cells)[:300]]; ws.column_dimensions[get_column_letter(idx)].width=min(max(max(map(len,vals),default=8)+2,11),42)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,(int,float)): cell.number_format="#,##0.00"
        ws=writer.book["По дням"]; chart=LineChart(); chart.title="Динамика выручки"; chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=len(bd)+1),titles_from_data=True); chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=len(bd)+1)); chart.height=7;chart.width=14;ws.add_chart(chart,"F2")
        ws=writer.book["По товарам"]; chart=BarChart(); chart.title="Выручка по товарам"; limit=min(len(bp)+1,11); chart.add_data(Reference(ws,min_col=3,min_row=1,max_row=limit),titles_from_data=True); chart.set_categories(Reference(ws,min_col=2,min_row=2,max_row=limit)); chart.height=7;chart.width=14;ws.add_chart(chart,"H2")
    return output.getvalue()
